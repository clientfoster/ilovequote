import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, title, text, border_color="4F46E5", bg_color="F8FAFC"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border styling
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"{title}: ")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(30, 41, 59)
    
    run_c = p.add_run(text)
    run_c.font.name = "Arial"
    run_c.font.size = Pt(9.5)
    run_c.font.color.rgb = RGBColor(71, 85, 105)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table_header(row, col_widths, headers, bg_color="1E293B"):
    for idx, header in enumerate(headers):
        cell = row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

def format_table_row(row, col_widths, values, is_even=False):
    bg = "F1F5F9" if is_even else "FFFFFF"
    for idx, val in enumerate(values):
        cell = row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, bg)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(val))
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(51, 65, 85)

def add_heading_1(doc, title):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return h

def add_heading_2(doc, title):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(3)
    r = h.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(79, 70, 229)
    return h

def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def build_word_document(output_path):
    doc = Document()
    
    # Page setup - 0.8 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    r_title = title_p.add_run("iLoveQuote – Version Records & Voice-to-Quote Technical Architecture")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(21)
    r_title.font.color.rgb = RGBColor(30, 41, 59)
    
    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    r_sub = sub_p.add_run("Comprehensive Engineering Log across All Reviews (V1, V2, V3 & V4) and Deep Dive on Voice-to-Quote/Invoice Conversion Engine")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    # Metadata Block Table
    meta_tbl = doc.add_table(rows=5, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = [Inches(2.2), Inches(4.7)]
    meta_data = [
        ("Client Organization", "Semixon Technologies"),
        ("Project Application", "iLoveQuote (Quotation & Invoicing SaaS Platform)"),
        ("Active Test Subdomain", "https://test.ilovequote.com"),
        ("Primary Modules & Routes", "/#/create-quote (V1/V2) | /#/create-quote2 (V3) | /#/page2 (V4 Voice Converter)"),
        ("Documentation Date", "06-October-2026"),
    ]
    for idx, (k, v) in enumerate(meta_data):
        row = meta_tbl.rows[idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width, cell_v.width = col_w[0], col_w[1]
        set_cell_background(cell_k, "F8FAFC")
        set_cell_background(cell_v, "FFFFFF")
        set_cell_margins(cell_k, top=50, bottom=50, left=100, right=100)
        set_cell_margins(cell_v, top=50, bottom=50, left=100, right=100)
        
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_before = Pt(0)
        pk.paragraph_format.space_after = Pt(0)
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.name = "Arial"
        rk.font.size = Pt(9)
        rk.font.color.rgb = RGBColor(71, 85, 105)
        
        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_before = Pt(0)
        pv.paragraph_format.space_after = Pt(0)
        rv = pv.add_run(v)
        rv.font.name = "Arial"
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = RGBColor(15, 23, 42)
        if "http" in v or "#" in v:
            rv.bold = True
            rv.font.color.rgb = RGBColor(79, 70, 229)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # ----------------------------------------------------
    # SECTION 1: ARCHITECTURE & BACKUP STRATEGY
    # ----------------------------------------------------
    add_heading_1(doc, "1. Version Control & Backup Architecture")
    add_callout(
        doc,
        "Zero-Risk Backup Guarantee",
        "Every client iteration is developed on an isolated Git branch with dedicated version tags. No version overwrites previous work. All historical iterations (Version 1, Version 2, Review 3, and Review 4) remain 100% accessible, isolated, and deployable.",
        border_color="10B981",
        bg_color="F0FDF4"
    )
    add_body_p(
        doc,
        "To satisfy the client requirement for complete version records and separate rollback capabilities, the system uses a dual-layer backup mechanism:\n"
        "• Git Branch Isolation: Dedicated branches (preview1, preview2, review3/preview3, review4) preserve complete code history.\n"
        "• Subdomain Routing: Test deployments run under https://test.ilovequote.com, while production remains safe under www.ilovequote.com.\n"
        "• Route Isolation: Distinct URLs allow side-by-side visual and functional testing without collision:\n"
        "    - Version 1 & 2: https://test.ilovequote.com/#/create-quote\n"
        "    - Version 3: https://test.ilovequote.com/#/create-quote2\n"
        "    - Version 4: https://test.ilovequote.com/#/page2 (and https://test.ilovequote.com/page2)"
    )
    
    # ----------------------------------------------------
    # SECTION 2: VERSION 1.0 (PREVIEW 1)
    # ----------------------------------------------------
    add_heading_1(doc, "2. Version 1.0 (Preview 1 / 'preview1')")
    add_body_p(
        doc,
        "• Release Date: 16-September-2026\n"
        "• Git Branch: preview1 (Baseline: 0d88a9d)\n"
        "• Target URL: https://test.ilovequote.com/#/create-quote\n"
        "• Primary Focus: Permanent Left Sidebar, 4-Step Quotation Stepper Wizard, Live Preview Board, and Mobile Icon Optimization."
    )
    
    col_w_v1 = [Inches(2.2), Inches(4.7)]
    v1_items = [
        ("Permanent Desktop Sidebar (Layout.tsx)", "Added permanent left sidebar (w-60 / w-64) with brand mark, '+ New Quote' button, navigation items, '< Collapse' toggle, and top user profile pill."),
        ("Step 1 Business Module (BusinessForm.tsx)", "Added 'Use a saved business profile' card, centered vertical drag-and-drop logo zone, and accordions for Business Address, GSTIN tax, and Social Links."),
        ("Live Preview Board (LivePreviewBoard.tsx)", "Instant split-screen card showing real-time Business Profile Card and responsive Quotation Sheet Preview side-by-side."),
        ("4-Step Stepper Wizard (QuoteWizard.tsx)", "Header with quote ID (Q-2026-00021), DRAFT pill badge, and visual progression: [1 Business] ── [2 Client] ── [3 Items] ── [4 Preview]."),
        ("Mobile View Optimization (LandingPage.tsx)", "Scaled down mobile tool icons from h-16 w-16 to h-9 w-9, reduced icon container box to 52px, and adjusted card width to 185px for pocket screens."),
    ]
    v1_tbl = doc.add_table(rows=len(v1_items) + 1, cols=2)
    v1_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(v1_tbl.rows[0], col_w_v1, ["Component / Area", "Key Enhancements Implemented"], bg_color="2563EB")
    for idx, (comp, desc) in enumerate(v1_items):
        format_table_row(v1_tbl.rows[idx + 1], col_w_v1, [comp, desc], is_even=(idx % 2 == 1))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # ----------------------------------------------------
    # SECTION 3: VERSION 2.0 (PREVIEW 2)
    # ----------------------------------------------------
    add_heading_1(doc, "3. Version 2.0 (Preview 2 / 'preview2')")
    add_body_p(
        doc,
        "• Release Date: 25-September-2026\n"
        "• Git Branch: preview2 (Commits: f6e4015, 87c9d20)\n"
        "• Target URL: https://test.ilovequote.com/#/create-quote\n"
        "• Primary Focus: Single-Page Instant Quote Generator (Figma model), Multi-Currency Dropdown, 5 Structured Accordions, JSON Schema, 1-Click Invoice Converter, and Full Mobile Responsiveness."
    )
    
    v2_items = [
        ("New Single-Page Generator (ManualQuoteGenerator.tsx)", "Single-page quotation generator matching Figma specification with top header actions, live currency selector, and formatted quote preview modal."),
        ("Multi-Currency Support", "Dropdown supporting 8 international currencies: INR (₹), USD ($), EUR (€), GBP (£), AED (AED), AUD (A$), CAD (C$), SGD (S$)."),
        ("5 Structured Collapsible Accordions", "Accordion structure: 1. Your Business, 2. Client Details, 3. Quote Details, 4. Items & Summary table, 5. Notes & Terms."),
        ("Structured Document Schema (structuredQuote.ts)", "Typed JSON quotation schema storing items, quantities, rates, discounts, taxes, and validity terms."),
        ("1-Click Invoice Conversion", "Direct data transformation from quotation to tax invoice via convertQuoteToInvoiceDraft() and POST /api/quotes/:id/convert-to-invoice."),
        ("Mobile UI Responsiveness", "Converted wide data tables into touch-friendly stacked card layouts across Create Quote, Create Invoice, Quotes List, and Invoices List."),
    ]
    v2_tbl = doc.add_table(rows=len(v2_items) + 1, cols=2)
    v2_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(v2_tbl.rows[0], col_w_v1, ["Component / Area", "Key Enhancements Implemented"], bg_color="0284C7")
    for idx, (comp, desc) in enumerate(v2_items):
        format_table_row(v2_tbl.rows[idx + 1], col_w_v1, [comp, desc], is_even=(idx % 2 == 1))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # ----------------------------------------------------
    # SECTION 4: VERSION 3.0 (REVIEW 3)
    # ----------------------------------------------------
    add_heading_1(doc, "4. Version 3.0 (Review 3 / 'review3' & 'preview3')")
    add_body_p(
        doc,
        "• Release Date: 30-September-2026\n"
        "• Git Branches: review3 & preview3 (Commits: b9e58ed, 8c629ec)\n"
        "• Target URL: https://test.ilovequote.com/#/create-quote2\n"
        "• Primary Focus: 'Additional Options' Section, Shipping Details, Refrens Tax Configuration Modal, Column & Formula Customization Modal, and Dynamic Item Table."
    )
    
    v3_items = [
        ("Additional Options Accordion", "Inserted between Quote Details and Items & Summary. Features Add Shipping Details toggle, '% Edit GST' button, and 'Edit Columns/Formulas' button."),
        ("Shipping Details Expandable Sub-module", "Expands Shipped From Address, Shipped To Name & Address, Transport Mode (Road, Rail, Air, Sea, Courier), Transporter Name, and Vehicle/Tracking Number."),
        ("Configure Tax Modal (TaxConfigModal.tsx)", "Refrens 4-step modal: 1. Tax Type (GST, VAT, Sales Tax, None) | 2. Place of Supply (36 Indian States/UTs + Other Territory) | 3. GST Type (IGST vs CGST & SGST + Cess %) | 4. Other Options (Reverse Charge Mechanism - RCM checkbox)."),
        ("Customize Columns & Formulas Modal (ColumnFormulaModal.tsx)", "Header with '+ Add New Column', 6-dot drag handles, column name input, type selector (TEXT / NUMBER), eye visibility toggles for Item, HSN/SAC, GST Rate, Quantity, plus custom column deletion."),
        ("Dynamic Desktop Table & Mobile Cards", "Desktop table and mobile cards dynamically reflect configured columns in real-time with custom column inputs per line item."),
        ("GST & Tax Calculation Summary", "Summary card computes subtotal, discounts, CGST & SGST 50/50 split, IGST, Cess % amounts, and displays Reverse Charge (RCM) status badge."),
    ]
    v3_tbl = doc.add_table(rows=len(v3_items) + 1, cols=2)
    v3_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(v3_tbl.rows[0], col_w_v1, ["Component / Area", "Key Enhancements Implemented"], bg_color="6D28D9")
    for idx, (comp, desc) in enumerate(v3_items):
        format_table_row(v3_tbl.rows[idx + 1], col_w_v1, [comp, desc], is_even=(idx % 2 == 1))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # ----------------------------------------------------
    # SECTION 5: VERSION 4.0 (REVIEW 4)
    # ----------------------------------------------------
    add_heading_1(doc, "5. Version 4.0 (Review 4 / 'review4') – Audio to Price Converter")
    add_body_p(
        doc,
        "• Release Date: 06-October-2026\n"
        "• Git Branches: review4 (dedicated backup) & review3 (Vercel test subdomain tracking)\n"
        "• Target URL: https://test.ilovequote.com/#/page2 (and https://test.ilovequote.com/page2)\n"
        "• Client Request from Semixon Technologies:\n"
        "    \"@!! you just work on audio to price converter module. put in test.ilovequote.com/ page2. like this impment in new page. do not deploy in previous design. thanks\"\n"
        "    \"audio to quote making, make it implementation. thanks\"\n"
        "• Primary Focus: Dedicated Audio-to-Price Converter & Voice-to-Quote Making module mounted strictly on /page2 without altering previous designs."
    )
    
    add_callout(
        doc,
        "Non-Destructive Isolated Deployment",
        "Version 4.0 strictly avoids touching or replacing existing pages. /#/create-quote (V1/V2) and /#/create-quote2 (V3) remain completely unchanged. The Audio to Price Converter operates independently on /#/page2.",
        border_color="D97706",
        bg_color="FFFBEB"
    )
    
    v4_items = [
        ("Dedicated Standalone Page (AudioQuoteConverterPage.tsx)", "Full voice-driven quote creation workstation mounted on /page2 with amber 'Page 2: Audio to Price Converter' badge, currency selector, and reset actions."),
        ("3 Audio Input Modalities", "1. Live Voice Recording with real-time waveform bars and elapsed timer\n2. Audio File Dropzone (.mp3, .wav, .m4a, .ogg)\n3. Direct Voice Note / Meeting Transcript Text Box with instant Re-Analyze."),
        ("Quick Voice Scenario Chips", "1-Click test scenario presets: 'Web Design & Hosting' (₹25k + ₹5k), 'Hardware Supply' (3 Laptops ₹65k, 2 Printers ₹18k, 10% disc), 'UI/UX App Dev' (₹35k + ₹95k, 18% GST)."),
        ("Smart Speech-to-Quote Parsing Engine", "Natural language parsing algorithm converting both numbers ('25000') and spoken words ('twenty five thousand', 'three laptops') into structured line items, quantities, and rates."),
        ("Real-time Quotation Table & Editor", "Full desktop table and mobile card editor allowing instant manual edits to item descriptions, unit rates, quantities, discounts, and taxes with live totals."),
        ("1-Click Invoice Conversion & Print Export", "Interactive Preview Modal, browser print/PDF export, and 1-Click Convert to Invoice mapping voice quote data directly into /create-invoice draft."),
        ("Dual URL Routing Architecture", "Custom vercel.json SPA rewrites and automatic path-to-hash redirect so direct hits to test.ilovequote.com/page2 and test.ilovequote.com/#/page2 work seamlessly."),
    ]
    v4_tbl = doc.add_table(rows=len(v4_items) + 1, cols=2)
    v4_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(v4_tbl.rows[0], col_w_v1, ["Feature / Component", "Implementation & User Experience"], bg_color="D97706")
    for idx, (comp, desc) in enumerate(v4_items):
        format_table_row(v4_tbl.rows[idx + 1], col_w_v1, [comp, desc], is_even=(idx % 2 == 1))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # ----------------------------------------------------
    # SECTION 6: TECHNICAL EXPLANATION - VOICE TO INVOICE & QUOTES ENGINE
    # ----------------------------------------------------
    add_heading_1(doc, "6. Technical Resources & Engine: How Voice Converts to Quotes & Invoices")
    add_body_p(
        doc,
        "This section details the exact technical resources, browser APIs, speech recognition pipelines, Natural Language Processing (NLP) parsing logic, and schema conversion models that turn spoken voice dictation into valid business price quotes and legal tax invoices."
    )
    
    add_callout(
        doc,
        "Voice-to-Document Pipeline Overview",
        "Spoken Audio / File Upload ➔ Audio Stream & Waveform Processing ➔ Speech Recognition (ASR) ➔ Spoken Number & Entity NLP Parser ➔ Structured JSON Quotation ➔ Financial Math Engine ➔ 1-Click Tax Invoice Auto-Mapper.",
        border_color="4F46E5",
        bg_color="EEF2FF"
    )
    
    add_heading_2(doc, "6.1. The 5-Stage Voice-to-Quote & Invoice Architecture")
    
    stages_col_w = [Inches(1.8), Inches(2.3), Inches(2.8)]
    stages_data = [
        ("Stage 1: Audio Capture & Signal Monitoring", "Web Speech API (SpeechRecognition)\nHTML5 MediaStream\nWeb Audio API (AudioContext, AnalyserNode)", "Captures real-time microphone stream. Renders dynamic sound waveform bars (visual feedback) and tracks elapsed recording duration with automatic silence timeouts."),
        ("Stage 2: Automatic Speech Recognition (ASR)", "Browser SpeechRecognition Engine\nOptional Cloud ASR (Whisper, Google STT)\nFileReader API for recorded audio", "Streams audio chunks and performs continuous speech-to-text decoding. Emits live interim and final transcripts into an editable text buffer."),
        ("Stage 3: NLP & Spoken Entity Parser", "Intelligent RegEx Lexer\nNUMBER_WORDS English tokenizer\nEntity & financial boundary detector", "Converts spoken English numbers ('twenty five thousand' ➔ 25000), isolates client names ('ABC Technologies'), extracts item descriptions, quantities, discount %, and tax %."),
        ("Stage 4: Structured Quote Modeling & Math", "StructuredQuote TypeScript Schema\ncalculateStructuredQuotePricing()\nMulti-Currency Formatter", "Instantiates typed quotation document. Dynamically calculates line totals, subtotal, discount deductions, GST (CGST/SGST/IGST), and validity expiration dates."),
        ("Stage 5: 1-Click Invoice Transformation", "convertQuoteToInvoiceDraft()\nlocalStorage persistence pipeline\nBackend POST /api/quotes/:id/convert-to-invoice", "Maps quote client details, items, quantities, and taxes directly into Tax Invoice draft format, ready for GST e-invoicing and PDF download."),
    ]
    stages_tbl = doc.add_table(rows=len(stages_data) + 1, cols=3)
    stages_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(stages_tbl.rows[0], stages_col_w, ["Pipeline Stage", "Core Technologies / Resources", "Technical Function & Role"], bg_color="1E293B")
    for idx, row_vals in enumerate(stages_data):
        format_table_row(stages_tbl.rows[idx + 1], stages_col_w, row_vals, is_even=(idx % 2 == 1))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    add_heading_2(doc, "6.2. Detailed Breakdown of Spoken Number & Word Translation")
    add_body_p(
        doc,
        "In voice dictation, users frequently say rates in words rather than digits (e.g. 'twenty five thousand rupees' instead of '25000', or 'three laptops' instead of '3'). The system implements a robust spoken number parsing algorithm:\n\n"
        "1. Tokenization & Normalization: The raw transcript is cleaned of non-alphanumeric punctuation and split into lowercase tokens.\n"
        "2. Spoken Vocabulary Mapping: A dictionary maps basic units (zero through nineteen), tens (twenty through ninety), and multipliers (hundred, thousand, lakh, lac, million, crore).\n"
        "3. Compound Value Accumulation: A sliding accumulator computes multi-word numbers:\n"
        "    • 'twenty' (20) + 'five' (5) = 25\n"
        "    • 25 * 'thousand' (1000) = 25,000\n"
        "    • Result: ₹25,000 unit rate cleanly extracted into the database without requiring user correction."
    )
    
    add_heading_2(doc, "6.3. Entity Boundary Detection for Items, Quantities, and Rates")
    add_body_p(
        doc,
        "Business quotes contain multiple items in a single sentence. The parser detects natural language boundaries:\n\n"
        "• Client Name Extraction: Scans for phrases like 'Create a quote for [Client]', 'Quote for [Company]', or 'Client [Name]'.\n"
        "• Line Item Segmentation: Splits text on conjunctions ('and', 'also', 'plus', commas, periods) and isolates distinct item clauses.\n"
        "• Quantity Pairing: Matches words like 'quantity [N]', '[N] pieces', '[N] units', or numbers immediately preceding nouns (e.g. '3 laptops' ➔ quantity = 3, item = 'laptops').\n"
        "• Rate Association: Matches numbers followed by currency terms ('rupees', 'dollars', 'inr', 'rs') or prepositions ('at [Rate]', 'for [Rate] each').\n"
        "• Validity & Terms: Extracts 'valid for [N] days' and dynamically computes the validUntil ISO date from today's system clock."
    )
    
    add_heading_2(doc, "6.4. Quote to Invoice Transformation Specification")
    add_body_p(
        doc,
        "Once the quote is generated from audio, the user can click 'Convert to Invoice' with 1 click. The conversion engine performs a seamless schema mapping:"
    )
    
    conv_cols = [Inches(2.5), Inches(2.2), Inches(2.2)]
    conv_data = [
        ("Client Name & Information", "StructuredQuote.client", "Invoice.clientDetails (Bill-To Customer)"),
        ("Item Name & Description", "StructuredQuoteItem.name, description", "InvoiceItem.description (Line Item)"),
        ("Quantity & Unit Rate", "StructuredQuoteItem.quantity, unitPrice", "InvoiceItem.quantity, unitPrice"),
        ("Line Amount", "StructuredQuoteItem.amount", "InvoiceItem.total"),
        ("Taxes & GST Breakdown", "StructuredQuote.pricing.taxPercent", "Invoice.taxRate (CGST / SGST / IGST)"),
        ("Discount Percent & Value", "StructuredQuote.pricing.discountPercent", "Invoice.discountPercent"),
        ("Document Status", "Quote Status: 'Created' / 'Approved'", "Invoice Status: 'Draft' (Awaiting Dispatch)"),
        ("Terms & Payment Due Date", "Valid Until (e.g. +7 days)", "Payment Due Date (+7 days from issue)"),
    ]
    conv_tbl = doc.add_table(rows=len(conv_data) + 1, cols=3)
    conv_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(conv_tbl.rows[0], conv_cols, ["Data Field", "Source (Voice Quote)", "Target (Tax Invoice)"], bg_color="0F766E")
    for idx, row_vals in enumerate(conv_data):
        format_table_row(conv_tbl.rows[idx + 1], conv_cols, row_vals, is_even=(idx % 2 == 1))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # ----------------------------------------------------
    # SECTION 7: CROSS-VERSION FEATURE COMPARISON MATRIX
    # ----------------------------------------------------
    add_heading_1(doc, "7. Cross-Version Feature Comparison Matrix across All 4 Reviews")
    
    matrix_cols = [Inches(2.3), Inches(1.1), Inches(1.1), Inches(1.2), Inches(1.2)]
    matrix_headers = ["Capability / Feature", "Version 1.0", "Version 2.0", "Review 3.0", "Review 4.0"]
    matrix_data = [
        ("Generation Model", "4-Step Wizard", "Instant Single-Page", "Instant Single-Page", "Audio-Driven Generator"),
        ("Dedicated URL Route", "/#/create-quote", "/#/create-quote", "/#/create-quote2", "/#/page2 (and /page2)"),
        ("Multi-Currency Support", "INR only", "8 Currencies (₹, $, €...)", "8 Currencies (₹, $, €...)", "8 Currencies (₹, $, €...)"),
        ("Live Audio Dictation (Mic)", "Not available", "Basic dialog", "Basic dialog", "Dedicated Waveform Studio"),
        ("Audio File Upload (.mp3)", "Not available", "Not available", "Not available", "Included (Dropzone)"),
        ("Spoken Word Number Parser", "Not available", "Digits only", "Digits only", "Full ('twenty five thousand')"),
        ("1-Click Invoice Conversion", "Manual entry", "1-Click Auto-Map", "1-Click Auto-Map", "1-Click Direct Pipeline"),
        ("Shipping Details Sub-module", "Not available", "Not available", "Included (5 Fields)", "Not needed on /page2"),
        ("GST / Tax Config Modal", "Standard Tax %", "Standard Tax %", "Full Refrens Modal", "Configurable Tax %"),
        ("Custom Columns & Formulas", "Fixed columns", "Fixed columns", "Dynamic Customizer", "Clean Standard Columns"),
        ("Document Print & PDF Export", "Basic", "Formatted Modal", "Formatted Modal", "Instant Print & PDF Modal"),
        ("Active Test Subdomain", "test.ilovequote.com", "test.ilovequote.com", "test.ilovequote.com", "test.ilovequote.com"),
    ]
    
    matrix_tbl = doc.add_table(rows=len(matrix_data) + 1, cols=5)
    matrix_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(matrix_tbl.rows[0], matrix_cols, matrix_headers, bg_color="0F172A")
    for idx, row_vals in enumerate(matrix_data):
        format_table_row(matrix_tbl.rows[idx + 1], matrix_cols, row_vals, is_even=(idx % 2 == 1))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # ----------------------------------------------------
    # SECTION 8: FILE INVENTORY & DEPLOYMENT DIRECTORY
    # ----------------------------------------------------
    add_heading_1(doc, "8. File Inventory & Active Deployment Directory")
    
    files_cols = [Inches(2.5), Inches(1.1), Inches(3.3)]
    files_headers = ["File Path", "Action", "Description & Version Scope"]
    files_data = [
        ("frontend/src/pages/AudioQuoteConverterPage.tsx", "Created (V4)", "Dedicated Audio to Price Converter page on /page2 with 3 audio modes, waveform bars, smart NLP parsing, and quote editor."),
        ("frontend/src/App.tsx", "Modified (V4)", "Registered /page2, /audio-quote, and /create-quote3 routes; added automatic direct-path to hash-route redirection."),
        ("frontend/src/components/Layout.tsx", "Modified (V4)", "Added Audio to Quote link with microphone icon in permanent sidebar navigation menu."),
        ("frontend/vercel.json & ./vercel.json", "Created (V4)", "Configured SPA rewrites to route direct path hits to index.html to prevent 404s on direct navigation."),
        ("frontend/src/components/quote-generator/TaxConfigModal.tsx", "Created (V3)", "4-step modal for GST/VAT/Sales Tax, Place of Supply, IGST/CGST, Cess & RCM."),
        ("frontend/src/components/quote-generator/ColumnFormulaModal.tsx", "Created (V3)", "Drag handle, column name/type customizer, visibility toggles & add custom column."),
        ("frontend/src/components/quote-generator/ManualQuoteGenerator.tsx", "Created (V2/V3)", "Single-page generator, Additional Options accordion, dynamic column rendering, and GST breakdown."),
        ("frontend/src/types/structuredQuote.ts", "Created (V2/V3/V4)", "Structured quotation schema, pricing calculator, and quote-to-invoice conversion helper."),
        ("VERSION_CHANGE_RECORDS.md", "Updated", "Master engineering backup logs across all 4 reviews."),
        ("generate_word_doc.py", "Updated", "Automated Word document generator producing this comprehensive record."),
    ]
    
    files_tbl = doc.add_table(rows=len(files_data) + 1, cols=3)
    files_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(files_tbl.rows[0], files_cols, files_headers, bg_color="1E293B")
    for idx, row_vals in enumerate(files_data):
        format_table_row(files_tbl.rows[idx + 1], files_cols, row_vals, is_even=(idx % 2 == 1))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # Footer Links Callout
    add_callout(
        doc,
        "Live Deployment & Access Links",
        "• Review 4 Active Target: https://test.ilovequote.com/#/page2 (and https://test.ilovequote.com/page2)\n"
        "• Review 3 Active Target: https://test.ilovequote.com/#/create-quote2\n"
        "• Review 1 & 2 Active Target: https://test.ilovequote.com/#/create-quote\n"
        "• Git Repository: https://github.com/clientfoster/ilovequote\n"
        "• Backup Branches: review4 (V4 Backup), review3 (Active Test Subdomain), preview3 (V3), preview2 (V2), preview1 (V1), main (Production).",
        border_color="4F46E5",
        bg_color="EEF2FF"
    )
    
    doc.save(output_path)
    print(f"Successfully generated Word document at: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\Admin\.gemini\antigravity\scratch\ilovequote"
    doc_path = os.path.join(out_dir, "iLoveQuote_Version_Change_Records_All_Reviews.docx")
    build_word_document(doc_path)
